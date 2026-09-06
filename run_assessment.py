"""Run the reproducible rule-based assessment and write the result workbook."""

import argparse
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.table import TableColumn

from claim_assessment import assess_reports
from evidence_assessment import analyse_evidence
from report_data import AssessmentReport, load_all_reports


SUMMARY_HEADERS = [
    "assessment_group", "total_claims", "manual_assessed", "matching", "mismatching", "agreement_rate",
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
ALT_FILL = PatternFill("solid", fgColor="D9EAF7")
HEADER_FONT = Font(color="FFFFFF", bold=True)
THIN_BORDER = Border(
    left=Side(style="thin", color="B7C9D6"),
    right=Side(style="thin", color="B7C9D6"),
    top=Side(style="thin", color="B7C9D6"),
    bottom=Side(style="thin", color="B7C9D6"),
)


# Read one worksheet as dictionaries keyed by its first-row column headers.
def read_rows(workbook, sheet_name: str) -> list[dict[str, object]]:
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Workbook is missing required sheet: {sheet_name}")
    values = list(workbook[sheet_name].values)
    headers = [str(value or "").strip() for value in values[0]] if values else []
    return [
        {headers[index]: row[index] if index < len(row) else None for index in range(len(headers))}
        for row in values[1:]
        if any(value not in (None, "") for value in row)
    ]


# Normalise user-entered labels before exact agreement comparisons.
def normalise_label(value: object) -> str:
    return str(value or "").strip().casefold()


# Attach manual labels after automatic scoring; manual labels never influence scoring.
def attach_manual_comparison(
    assessments: list[dict[str, object]], reports: list[AssessmentReport],
) -> list[dict[str, object]]:
    manual_lookup = {
        report.case_id: normalise_label(report.manual_assessment)
        for report in reports
    }
    output: list[dict[str, object]] = []
    for assessment in assessments:
        row = dict(assessment)
        manual = manual_lookup.get(str(row["claim_id"]), "")
        row["manual_assessment"] = manual
        if not manual:
            row["comparison_result"] = "not_reviewed"
        elif manual == normalise_label(row["assessment"]):
            row["comparison_result"] = "match"
        else:
            row["comparison_result"] = "mismatch"
        output.append(row)
    return output


# Calculate agreement only for claims that have a manual reference assessment.
def agreement_metrics(rows: list[dict[str, object]]) -> dict[str, object]:
    assessed = [row for row in rows if row["comparison_result"] != "not_reviewed"]
    matching = sum(row["comparison_result"] == "match" for row in assessed)
    manual_assessed = len(assessed)
    return {
        "total_claims": len(rows),
        "manual_assessed": manual_assessed,
        "matching": matching,
        "mismatching": manual_assessed - matching,
        "agreement_rate": matching / manual_assessed if manual_assessed else None,
    }


# Produce the report, synthetic, and combined rows for the summary worksheet.
def agreement_summary(
    report_assessments: list[dict[str, object]], synthetic_assessments: list[dict[str, object]],
) -> list[dict[str, object]]:
    groups = [
        ("Report claims", report_assessments),
        ("Synthetic claims", synthetic_assessments),
        ("All claims", report_assessments + synthetic_assessments),
    ]
    return [{"assessment_group": name, **agreement_metrics(rows)} for name, rows in groups]


# Locate a required workbook column by header rather than relying on a fixed position.
def _header_column(worksheet, header: str) -> int:
    for column in range(1, worksheet.max_column + 1):
        if str(worksheet.cell(row=1, column=column).value or "").strip() == header:
            return column
    raise ValueError(f"Worksheet {worksheet.title} is missing required column: {header}")


# Extend the existing Excel table when the auto-assessment column is first added.
def _extend_table_range(worksheet, final_column: int) -> None:
    """Extend a table and synchronise its metadata with the header row."""
    for table in worksheet.tables.values():
        min_column, min_row, _, max_row = range_boundaries(table.ref)
        if min_row == 1 and min_column == 1:
            table.ref = f"A1:{get_column_letter(final_column)}{max_row}"
            table.tableColumns = [
                TableColumn(
                    id=column,
                    name=str(worksheet.cell(row=1, column=column).value or f"Column{column}"),
                )
                for column in range(1, final_column + 1)
            ]


# Add or refresh automatic assessments while copying the existing manual-column style.
def add_auto_assessment_column(
    worksheet, identifier_column: str, assessments: list[dict[str, object]],
) -> None:
    id_column = _header_column(worksheet, identifier_column)
    manual_column = _header_column(worksheet, "manual_assessment")
    assessment_by_id = {str(row["claim_id"]): row["assessment"] for row in assessments}

    try:
        auto_column = _header_column(worksheet, "auto_assessment")
    except ValueError:
        auto_column = worksheet.max_column + 1
        header_cell = worksheet.cell(row=1, column=auto_column)
        header_cell.value = "auto_assessment"
        header_cell._style = copy(worksheet.cell(row=1, column=manual_column)._style)
        header_cell.alignment = copy(worksheet.cell(row=1, column=manual_column).alignment)
        worksheet.column_dimensions[get_column_letter(auto_column)].width = worksheet.column_dimensions[
            get_column_letter(manual_column)
        ].width
        for row_number in range(2, worksheet.max_row + 1):
            worksheet.cell(row=row_number, column=auto_column)._style = copy(
                worksheet.cell(row=row_number, column=manual_column)._style
            )
        _extend_table_range(worksheet, auto_column)

    for row_number in range(2, worksheet.max_row + 1):
        identifier = str(worksheet.cell(row=row_number, column=id_column).value or "").strip()
        worksheet.cell(row=row_number, column=auto_column).value = assessment_by_id.get(identifier, "")


# Apply the compact table style used by the generated assessment summary.
def _style_summary(sheet, header_rows: set[int]) -> None:
    for row in sheet.iter_rows(min_row=3, max_row=sheet.max_row):
        for cell in row:
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if cell.row in header_rows:
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            elif cell.row % 2 == 0:
                cell.fill = ALT_FILL


# Recreate the summary sheet from the latest automatic and manual comparison results.
def add_assessment_summary(
    workbook, report_assessments: list[dict[str, object]], synthetic_assessments: list[dict[str, object]],
) -> None:
    if "Assessment Summary" in workbook.sheetnames:
        del workbook["Assessment Summary"]
    sheet = workbook.create_sheet("Assessment Summary")
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:F1")
    sheet["A1"] = "Assessment Summary"
    sheet["A1"].fill = HEADER_FILL
    sheet["A1"].font = Font(color="FFFFFF", bold=True, size=14)
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 28

    summary_rows = agreement_summary(report_assessments, synthetic_assessments)
    sheet.append([])
    sheet.append(SUMMARY_HEADERS)
    for row in summary_rows:
        sheet.append([row[header] for header in SUMMARY_HEADERS])
    _style_summary(sheet, {3})
    for row_number in range(4, 7):
        sheet.cell(row=row_number, column=6).number_format = "0.0%"
    for column, width in {"A": 24, "B": 20, "C": 18, "D": 16, "E": 16, "F": 16}.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A3"


# Copy the input workbook, update automatic assessments, and save the result workbook.
def write_result_workbook(
    input_path: Path,
    output_path: Path,
    report_assessments: list[dict[str, object]],
    synthetic_assessments: list[dict[str, object]],
) -> None:
    workbook = load_workbook(input_path)
    add_auto_assessment_column(workbook["Report Sources"], "report_id", report_assessments)
    add_auto_assessment_column(workbook["Synthetic Claims"], "synthetic_claim_id", synthetic_assessments)
    add_assessment_summary(workbook, report_assessments, synthetic_assessments)
    workbook.save(output_path)


# Parse command-line paths, run the rule engine, compare results, and write the workbook.
def main() -> None:
    parser = argparse.ArgumentParser(description="Run rule-based evidence assessment on the simplified workbook schema (no LLM).")
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()

    workbook = load_workbook(args.input, read_only=True, data_only=True)
    try:
        report_rows = read_rows(workbook, "Report Sources")
        synthetic_rows = read_rows(workbook, "Synthetic Claims")
        source_rows = read_rows(workbook, "Source Evidence Excerpts")
        report_objects, synthetic_objects = load_all_reports(
            report_rows, synthetic_rows, source_rows,
        )
        all_reports = report_objects + synthetic_objects
        for report in all_reports:
            analyse_evidence(report)
        assessments = assess_reports(all_reports)
    finally:
        workbook.close()

    report_assessments = attach_manual_comparison(
        [row for row in assessments if row["claim_source_type"] == "report"],
        report_objects,
    )
    synthetic_assessments = attach_manual_comparison(
        [row for row in assessments if row["claim_source_type"] == "synthetic"],
        synthetic_objects,
    )
    args.output.parent.mkdir(parents=True, exist_ok=True)
    write_result_workbook(args.input, args.output, report_assessments, synthetic_assessments)
    print(f"Result workbook: {args.output}")


if __name__ == "__main__":
    main()
